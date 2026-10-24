"""In-memory extraction for PDF/DOCX resumes."""
from __future__ import annotations

import re
from dataclasses import dataclass, field

import pymupdf
from docx import Document

from src.section_parser import Entry, parse_entries, split_sections
from src.utils import SKILL_TERMS


@dataclass
class ResumeData:
    """Structured fields and source text extracted from a resume."""
    name: str = ""
    email: str = ""
    phone: str = ""
    skills: list[str] = field(default_factory=list)
    education: list[Entry] = field(default_factory=list)
    projects: list[Entry] = field(default_factory=list)
    experience: list[Entry] = field(default_factory=list)
    certifications: list[str] = field(default_factory=list)
    profile_links: list[str] = field(default_factory=list)
    institutions: list[str] = field(default_factory=list)
    text: str = ""
    warnings: list[str] = field(default_factory=list)


def _pdf_text(data: bytes) -> tuple[str, list[str]]:
    try:
        with pymupdf.open(stream=data, filetype="pdf") as doc:
            pages = [page.get_text("text") for page in doc]
    except Exception as exc:
        raise ValueError("This PDF could not be opened. Check that it is a valid, unencrypted PDF.") from exc
    if not pages or not "".join(pages).strip():
        raise ValueError("No selectable text was found. This may be a scanned PDF; OCR is not included.")
    return "\n".join(pages), ([] if any(p.strip() for p in pages) else ["Some PDF pages contain no selectable text."])


def _docx_text(data: bytes) -> str:
    from io import BytesIO
    try:
        document = Document(BytesIO(data))
    except Exception as exc:
        raise ValueError("This DOCX could not be opened. Check that it is a valid Word document.") from exc
    pieces = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        pieces.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
    text = "\n".join(pieces)
    if not text.strip():
        raise ValueError("The document contains no readable text.")
    return text


def extract_resume(filename: str, data: bytes) -> ResumeData:
    """Extract a conservative set of resume fields. The input is never persisted."""
    if not data:
        raise ValueError("The uploaded file is empty.")
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix not in {"pdf", "docx"}:
        raise ValueError("Unsupported file type. Please upload a PDF or DOCX file.")
    text, warnings = _pdf_text(data) if suffix == "pdf" else (_docx_text(data), [])
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    email_match = re.search(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", text, re.I)
    phone_match = re.search(r"(?<!\w)(?:\+?[0-9][0-9\s().-]{7,}[0-9])(?!\w)", text)
    name = ""
    for line in lines[:5]:
        candidate = re.sub(r"\s+", " ", line).strip()
        if (len(candidate.split()) in (2, 3) and len(candidate) < 50 and
                not re.search(r"@|https?://|\d|resume|curriculum vitae", candidate, re.I)):
            name = candidate
            break
    sections = split_sections(text)
    full_lower = text.casefold()
    skills = [skill for skill in SKILL_TERMS if re.search(r"(?<!\w)" + re.escape(skill) + r"(?!\w)", full_lower)]
    skill_section = sections.get("skills", "")
    for token in re.split(r"[,;|\n•]+", skill_section):
        clean = token.strip(" .-–")
        if clean and len(clean) < 45 and clean.casefold() not in {s.casefold() for s in skills}:
            skills.append(clean)
    education = parse_entries(sections.get("education", ""))
    institutions = [entry.title for entry in education if re.search(r"university|college|institute|school", entry.title, re.I)]
    certs = [x.title for x in parse_entries(sections.get("certifications", ""))]
    projects = parse_entries(sections.get("projects", ""))
    experience = parse_entries(sections.get("experience", ""))
    links = re.findall(r"(?:https?://|www\.)\S+", text, re.I)
    if not skills: warnings.append("No skills could be confidently identified; check the preview and add missing context manually.")
    return ResumeData(name=name, email=email_match.group(0) if email_match else "",
                      phone=phone_match.group(0).strip() if phone_match else "", skills=skills,
                      education=education, projects=projects, experience=experience,
                      certifications=certs, profile_links=links, institutions=institutions,
                      text=text, warnings=warnings)
